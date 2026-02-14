#!/usr/bin/env python3
"""Remove all directional/advocacy language from Decision Matrix sheet."""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = load_workbook('/home/user/LB-Pharma-qsbs/LBRX_QSBS_Evecxia_Model.xlsx')
ws = wb['Decision Matrix']

thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)
header_font = Font(name='Calibri', bold=True, size=11)
light_yellow_fill = PatternFill(start_color='FFF9C4', end_color='FFF9C4', fill_type='solid')

# ============================================================
# 1. REMOVE "Advantage" column (D) from Head-to-Head Comparison
# ============================================================
# Row 35 header: remove "Advantage"
ws['D35'] = ''
ws['D35'].fill = PatternFill()  # clear fill

# Rows 36-43: clear all Advantage column values and formatting
for r in range(36, 44):
    ws.cell(row=r, column=4).value = None
    ws.cell(row=r, column=4).font = Font()
    ws.cell(row=r, column=4).fill = PatternFill()
    ws.cell(row=r, column=4).border = Border()

# Also fix any green-colored PATH 3 text in column D
# (already cleared above)

# ============================================================
# 2. REMOVE "VERDICT" section entirely (rows 45-48)
# ============================================================
for r in range(45, 49):
    for c in range(1, 7):
        cell = ws.cell(row=r, column=c)
        cell.value = None
        cell.font = Font()
        cell.fill = PatternFill()
        cell.border = Border()

# ============================================================
# 3. RENAME "Path 3 Favored?" to neutral language
# ============================================================
ws['A55'] = 'Path 3 EV > Path 2 EV?'
ws['A55'].font = header_font
ws['A55'].border = thin_border

# ============================================================
# 4. RENAME sensitivity section — remove "Path 3 Net EV" framing
# ============================================================
ws['A51'] = 'Expected net at varying P(Ph2 Success):'
ws['A51'].font = Font(italic=True, size=10)

# ============================================================
# 5. HEAD-TO-HEAD: remove "EV / Risk Ratio" row — it's designed to favor Path 3
# ============================================================
ws['A43'] = None
ws['B43'] = None
ws['C43'] = None
for c in range(1, 4):
    ws.cell(row=43, column=c).font = Font()
    ws.cell(row=43, column=c).border = Border()

# ============================================================
# 6. Relabel head-to-head section header to be neutral
# ============================================================
ws['A34'] = 'SIDE-BY-SIDE COMPARISON'

# ============================================================
# 7. Fix "Liquidity" row — was hardcoded "PATH 2" as advantage
# ============================================================
ws['D41'] = None  # remove hardcoded "TIED"
ws['D42'] = None  # remove hardcoded "PATH 2"

wb.save('/home/user/LB-Pharma-qsbs/LBRX_QSBS_Evecxia_Model.xlsx')
print("Decision Matrix: removed all directional language")
