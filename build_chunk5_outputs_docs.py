#!/usr/bin/env python3
"""
LBRX QSBS Evecxia Model - Chunk 5: Outputs Dashboard + README + Sources
Outputs: Key flags panel, scenario comparison, sensitivity table
README: Usage instructions, key tax rules, disclaimer
Sources: Full bibliography
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime

def build_outputs_tab(wb):
    """Build Outputs dashboard with key flags, scenario comparison, and sensitivity table"""
    ws = wb['Outputs']

    # Column widths
    ws.column_dimensions['A'].width = 40
    for col in ['B', 'C', 'D', 'E', 'F', 'G']:
        ws.column_dimensions[col].width = 18

    row = 1
    ws['A1'] = 'OUTPUTS DASHBOARD'
    ws['A1'].font = Font(name='Calibri', size=14, bold=True)
    row += 2

    # KEY FLAGS PANEL
    ws['A3'] = 'KEY FLAGS'
    ws['A3'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    ws['A3'].fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    ws.merge_cells('A3:C3')
    row = 4

    flags = [
        ('QSBS 5-year met at lock-up?', '=IF(lockup_date>=qsbs_5yr_date,"YES","NO")'),
        ('QSBS 5-year met at Phase 3 readout?', '=IF(ph3_date>=qsbs_5yr_date,"YES","NO")'),
        ('§1045 eligible?', '=IF(AND((lockup_date-serb_date)>180,Inputs!$B$33="YES"),"YES","NO")'),
        ('§1045 → §1202 chain achievable?', '=IF(AND(B6="YES",(evecxia_ph2_date-serb_date)>=1826),"YES","NO")'),
    ]

    for label, formula in flags:
        ws[f'A{row}'] = label
        ws[f'B{row}'] = formula
        ws[f'B{row}'].alignment = Alignment(horizontal='center')
        ws[f'B{row}'].font = Font(name='Calibri', size=11, bold=True)
        # Conditional formatting would go here (green for YES, red for NO)
        row += 1

    row += 2

    # SCENARIO COMPARISON TABLE
    ws[f'A{row}'] = 'SCENARIO COMPARISON'
    ws[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    ws[f'A{row}'].fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    ws.merge_cells(f'A{row}:G{row}')
    row += 1

    # Headers
    headers = ['Metric', 'S1: Lock-Up', 'S2: Ph3+', 'S3: Ph3-', 'S4: Dual Lot', 'S5: Eve+', 'S6: Eve-']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_num)
        cell.value = header
        cell.font = Font(name='Calibri', size=10, bold=True)
        cell.fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    row += 1

    # Pull key metrics from Scenarios tab
    comparison_metrics = [
        ('Gross Proceeds', ['=Scenarios!B9', '=Scenarios!C9', '=Scenarios!D9', '=Scenarios!D36', '=Scenarios!F9', '=Scenarios!G9'], '$#,##0'),
        ('Gain', ['=Scenarios!B11', '=Scenarios!C11', '=Scenarios!D11', '=Scenarios!D38', '=Scenarios!F11', '=Scenarios!G11'], '$#,##0'),
        ('§1202 Exclusion', ['=Scenarios!B18', '=Scenarios!C18', '=Scenarios!D18', '=Scenarios!D47', '=Scenarios!F18', '=Scenarios!G18'], '$#,##0'),
        ('Total Tax', ['=Scenarios!B24', '=Scenarios!C24', '=Scenarios!D24', '=Scenarios!D49', '=Scenarios!F24', '=Scenarios!G24'], '$#,##0'),
        ('Net After-Tax', ['=Scenarios!B26', '=Scenarios!C26', '=Scenarios!D26', '=Scenarios!D50', '=Scenarios!F26', '=Scenarios!G26'], '$#,##0'),
        ('MOIC', ['=Scenarios!B27', '=Scenarios!C27', '=Scenarios!D27', 'Varies', '=Scenarios!F27', '=Scenarios!G27'], '0.00x'),
    ]

    for metric_label, formulas, num_fmt in comparison_metrics:
        ws[f'A{row}'] = metric_label
        for col_num, formula in enumerate(formulas, 2):
            cell = ws.cell(row=row, column=col_num)
            cell.value = formula
            if num_fmt:
                cell.number_format = num_fmt
            cell.alignment = Alignment(horizontal='right')
        row += 1

    row += 2

    # SENSITIVITY TABLE (simplified 2-way)
    ws[f'A{row}'] = 'SENSITIVITY ANALYSIS: LBRX Price vs Tax Treatment'
    ws[f'A{row}'].font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    ws[f'A{row}'].fill = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
    ws.merge_cells(f'A{row}:G{row}')
    row += 1

    # Column headers (LBRX prices)
    prices = [5, 10, 15, 20, 30, 50, 100]
    ws[f'A{row}'] = 'Tax Treatment'
    for col_num, price in enumerate(prices, 2):
        cell = ws.cell(row=row, column=col_num)
        cell.value = f'${price}'
        cell.font = Font(name='Calibri', size=10, bold=True)
        cell.alignment = Alignment(horizontal='center')
    row += 1

    # Row: No QSBS (standard LTCG)
    ws[f'A{row}'] = 'No QSBS (standard LTCG)'
    for col_num, price in enumerate(prices, 2):
        cell = ws.cell(row=row, column=col_num)
        # Formula: proceeds - basis - tax
        cell.value = f'=serb_shares*{price}-(serb_shares*{price}-serb_basis)*combined_rate'
        cell.number_format = '$#,##0'
    row += 1

    # Row: QSBS §1202 Excluded
    ws[f'A{row}'] = 'QSBS §1202 Excluded'
    for col_num, price in enumerate(prices, 2):
        cell = ws.cell(row=row, column=col_num)
        # Formula: proceeds (no tax if gain < $10M)
        cell.value = f'=serb_shares*{price}-MAX(0,(serb_shares*{price}-serb_basis-10000000))*combined_rate'
        cell.number_format = '$#,##0'
    row += 1

    ws.freeze_panes = 'A4'


def build_readme_tab(wb):
    """Build README tab with usage instructions"""
    ws = wb['README']
    ws.column_dimensions['A'].width = 100

    content = [
        ('LBRX QSBS & Evecxia Scenario Model — README', 14, True),
        ('', 10, False),
        ('WHAT THIS MODEL DOES:', 12, True),
        ('This Excel model analyzes tax scenarios for LBRX (LB Pharmaceuticals) Series B/C investors under QSBS rules, including §1202 exclusion and §1045 rollover into Evecxia Therapeutics.', 10, False),
        ('', 10, False),
        ('HOW TO USE:', 12, True),
        ('1. Go to the Inputs tab', 10, True),
        ('2. Modify investor-specific assumptions in light blue cells:', 10, False),
        ('   - Federal LTCG rate, NIIT, state tax rate (rows 5-9)', 10, False),
        ('   - Series B cost basis and share price (rows 15-16)', 10, False),
        ('   - Lock-up price (row 42) — EDITABLE to test different sale prices', 10, False),
        ('   - Scenario valuations (rows 52-65) — Adjust based on your comp analysis', 10, False),
        ('3. Review Scenarios tab for 6 tax outcome scenarios', 10, False),
        ('4. Check Outputs dashboard for summary comparison', 10, False),
        ('', 10, False),
        ('KEY TAX RULES EMBEDDED IN THIS MODEL:', 12, True),
        ('• Dilution is not a taxable event; tax = sale proceeds minus basis, lot-by-lot', 10, False),
        ('• §1202 requires 5 full years of holding; March 2026 sale of May 2022 stock = no exclusion', 10, False),
        ('• Tranches are separate tax lots; Series B QSBS does not blend with Series C non-QSBS', 10, False),
        ('• IPO conversion to common does not erase lot identity', 10, False),
        ('• §1045 rollover can defer gain and preserve tacked holding period if requirements met', 10, False),
        ('', 10, False),
        ('IMPORTANT DISCLAIMER:', 12, True),
        ('This is an analytical model for illustrative purposes, NOT tax advice.', 10, True),
        ('Consult a qualified CPA or tax attorney before making investment decisions.', 10, True),
        ('§1045 rollovers have specific requirements; state tax treatment varies.', 10, True),
        ('', 10, False),
        ('For questions or issues with this model, contact the model creator.', 10, False),
        (f'Generated: {datetime.now().strftime("%Y-%m-%d %H:%M")}', 9, False),
    ]

    for row_num, (text, size, bold) in enumerate(content, 1):
        cell = ws[f'A{row_num}']
        cell.value = text
        cell.font = Font(name='Calibri', size=size, bold=bold)
        cell.alignment = Alignment(wrap_text=True, vertical='top')


def build_sources_tab(wb):
    """Build Sources tab with bibliography"""
    ws = wb['Sources']

    # Column widths
    ws.column_dimensions['A'].width = 5
    ws.column_dimensions['B'].width = 40
    ws.column_dimensions['C'].width = 70
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 15

    # Headers
    headers = ['#', 'Source', 'URL', 'Category', 'Date Accessed']
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')

    # Sources (from research)
    sources = [
        ('LBRX Current Stock Price', 'https://finance.yahoo.com/quote/LBRX/', 'Market Data', '2026-02-13'),
        ('LBRX IPO Announcement', 'https://www.cooley.com/news/coverage/2025/2025-09-12-lb-pharmaceuticals-announces-$285-million-upsized-ipo', 'Press Release', '2026-02-13'),
        ('LBRX Phase 2 ILLUMINATE-1', 'https://www.globenewswire.com/news-release/2026/01/26/3225570/0/en/LB-Pharmaceuticals-Initiates-Phase-2-ILLUMINATE-1-Trial-in-Bipolar-Depression-Expanding-LB-102-Development-Program.html', 'Clinical Data', '2026-02-13'),
        ('LBRX PIPE Financing Feb 2026', 'https://www.stocktitan.net/news/LBRX/lb-pharmaceuticals-announces-100-0-million-private-oct35p3986tv.html', 'Press Release', '2026-02-13'),
        ('Karuna BMY Acquisition', 'https://news.bms.com/news/details/2023/Bristol-Myers-Squibb-Strengthens-Neuroscience-Portfolio-with-Acquisition-of-Karuna-Therapeutics/default.aspx', 'Comp Valuation', '2026-02-13'),
        ('ITCI JNJ Acquisition', 'https://www.jnj.com/media-center/press-releases/johnson-johnson-strengthens-neuroscience-leadership-with-acquisition-of-intra-cellular-therapies-inc', 'Comp Valuation', '2026-02-13'),
        ('Cerevel AbbVie Acquisition', 'https://news.abbvie.com/2023-12-06-AbbVie-to-Acquire-Cerevel-Therapeutics-in-Transformative-Transaction-to-Strengthen-Neuroscience-Pipeline', 'Comp Valuation', '2026-02-13'),
        ('Cerevel Emraclidine Ph2 Failure', 'https://news.abbvie.com/2024-11-11-AbbVie-Provides-Update-on-Phase-2-Results-for-Emraclidine-in-Schizophrenia', 'Comp Valuation', '2026-02-13'),
        ('IRS §1202 QSBS Rules', 'https://www.irs.gov/instructions/i1040sd (Schedule D instructions)', 'Tax Authority', '2026-02-13'),
        ('IRS §1045 Rollover Rules', 'https://www.irs.gov/publications/p550 (Investment Income and Expenses)', 'Tax Authority', '2026-02-13'),
    ]

    for row_num, (source, url, category, date) in enumerate(sources, 2):
        ws[f'A{row_num}'] = row_num - 1
        ws[f'B{row_num}'] = source
        ws[f'C{row_num}'] = url
        ws[f'C{row_num}'].alignment = Alignment(wrap_text=True)
        ws[f'D{row_num}'] = category
        ws[f'E{row_num}'] = date

    ws.freeze_panes = 'A2'


def main():
    print("Building LBRX QSBS Evecxia Model - Chunk 5: Outputs + README + Sources...")

    # Load checkpoint
    wb = load_workbook('/home/user/LB-Pharma-qsbs/LBRX_QSBS_Evecxia_Model_v4_scenarios.xlsx')
    print("✓ Loaded checkpoint v4")

    # Build tabs
    build_outputs_tab(wb)
    print("✓ Built Outputs dashboard (flags, comparison, sensitivity)")

    build_readme_tab(wb)
    print("✓ Built README tab")

    build_sources_tab(wb)
    print("✓ Built Sources tab (10 key sources)")

    # Save checkpoint
    output_file = '/home/user/LB-Pharma-qsbs/LBRX_QSBS_Evecxia_Model_v5_complete.xlsx'
    wb.save(output_file)
    print(f"✓ Saved checkpoint: {output_file}")
    print("\nChunk 5 complete!")

if __name__ == '__main__':
    main()
