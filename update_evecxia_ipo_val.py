#!/usr/bin/env python3
"""Update Evecxia IPO valuation from $400M to $180M (Axsome IPO comp)"""
from openpyxl import load_workbook

wb = load_workbook('/home/user/LB-Pharma-qsbs/LBRX_QSBS_Evecxia_Model.xlsx')
ws = wb['Inputs']

# Update B64: Evecxia IPO Valuation from $400M to $180M
ws['B64'] = 180000000
ws['B64'].number_format = '$#,##0'

wb.save('/home/user/LB-Pharma-qsbs/LBRX_QSBS_Evecxia_Model.xlsx')
print("Updated Evecxia IPO Valuation: $400M -> $180M")
